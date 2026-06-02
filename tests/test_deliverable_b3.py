"""B-3 검수 — PDF/PPT 렌더 + 디자인 검수 + ★ 렌더물도 절대선 게이트.

핵심 증명(★):
  렌더된 PDF/PPT 도 guarded_send 를 거치며, 위반 시 real_fn_called=False
  그리고 rendered=False (렌더조차 미실행). 렌더물이 게이트를 우회하지 못함.

추가: L1/L2/L4 PDF/PPT 렌더 성공, 비전 디자인 검수 점수 산출, as_document 전달,
면책 자동주입이 렌더물에도 반영.
"""

import importlib.util
import sys
from pathlib import Path

import pytest

_REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_REPO))

from ai_pm import document_renderer as dr  # noqa: E402
from ai_pm.disclaimer_injector import REQUIRED_DISCLAIMER  # noqa: E402
from ai_pm.document_pipeline import render_and_deliver_document  # noqa: E402
from ai_pm.sample_data import sample_company  # noqa: E402
from ai_pm.sandbox import SYNTHETIC_MARK  # noqa: E402
from ai_pm.vision_review import DESIGN_RUBRIC, review_design  # noqa: E402

pytestmark = pytest.mark.skipif(
    not (dr.pptx_available() and dr.pdf_available()),
    reason="render libs (python-pptx/reportlab) not installed",
)


# ── 스킬 build 스크립트 동적 로드 ──
def _load(name: str):
    path = _REPO / "skills" / "vvs-deliverable-doc" / "scripts" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(f"_b3_{name}", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


BUILD_PPTX = _load("build_pptx")
BUILD_PDF = _load("build_pdf")


# ── as_document 전달 함수 스파이 (헤르메스 as_document 대역) ──
class DocSpy:
    def __init__(self):
        self.called = False
        self.deliveries = []

    def __call__(self, path, **meta):
        self.called = True
        self.deliveries.append({"path": path, **meta})
        return {"as_document": path, **meta}


from ai_pm.sample_data import sample_company  # noqa: E402


def _clean_data(**over):
    """원청 제출 수준 합성 회사 데이터(8섹션). 위반 주입은 over 로."""
    data = sample_company("L1")
    data.update(over)
    return data


def _good_claim(**over):
    base = {
        "synthetic": True,
        "marking": SYNTHETIC_MARK,
        "statement": "협력 실적 양호",
        "source": "[SYNTHETIC] internal",
    }
    base.update(over)
    return base


# ════════════════════════════════════════════════════════════════════
# 1) L1/L2/L4 PDF/PPT 렌더 성공 (게이트 통과 시)
# ════════════════════════════════════════════════════════════════════
@pytest.mark.parametrize("level", ["L1", "L2", "L4"])
def test_pptx_render_success_all_levels(level, tmp_path):
    spy = DocSpy()
    res = BUILD_PPTX.run(
        _clean_data(), spy, level=level, claims=[_good_claim()], out_dir=str(tmp_path)
    )
    assert res.delivered is True
    assert res.real_fn_called is True
    assert res.rendered is True
    assert res.render.fmt == "pptx"
    assert Path(res.render.path).exists()
    assert res.render.bytes > 0
    assert res.render.slide_count >= 2  # 표지 + 최소 1 섹션


@pytest.mark.parametrize("level", ["L1", "L2", "L4"])
def test_pdf_render_success_all_levels(level, tmp_path):
    spy = DocSpy()
    res = BUILD_PDF.run(
        _clean_data(), spy, level=level, claims=[_good_claim()], out_dir=str(tmp_path)
    )
    assert res.delivered is True and res.real_fn_called is True and res.rendered is True
    assert res.render.fmt == "pdf"
    assert Path(res.render.path).exists() and res.render.bytes > 0


# ════════════════════════════════════════════════════════════════════
# ★ 2) 렌더물도 절대선 게이트: 위반 시 real_fn_called=False + rendered=False
# ════════════════════════════════════════════════════════════════════
def test_pptx_banned_blocks_render_and_delivery():
    spy = DocSpy()
    res = BUILD_PPTX.run(_clean_data(company="저희가 무조건 업계 1위입니다"), spy)
    assert res.delivered is False
    assert res.real_fn_called is False
    assert res.rendered is False          # ← 렌더조차 미실행
    assert spy.called is False            # ← as_document 미전달 (렌더물도)
    assert any(v["kind"] == "banned_phrase" for v in res.violations)


def test_pdf_unverified_fact_blocks_render_and_delivery():
    spy = DocSpy()
    res = BUILD_PDF.run(_clean_data(), spy, claims=[_good_claim(source="")])
    assert res.real_fn_called is False and res.rendered is False
    assert spy.called is False
    assert any(v["kind"] == "fact" for v in res.violations)


def test_all_levels_formats_block_on_violation(tmp_path):
    """L1/L2/L4 × pptx/pdf 전부 위반 시 미렌더·미전달."""
    for level in ("L1", "L2", "L4"):
        for builder in (BUILD_PPTX, BUILD_PDF):
            spy = DocSpy()
            res = builder.run(
                _clean_data(company="확실히 100% 보장 무조건"),
                spy, level=level, out_dir=str(tmp_path),
            )
            assert res.real_fn_called is False
            assert res.rendered is False
            assert spy.called is False


# ════════════════════════════════════════════════════════════════════
# 3) 면책 자동주입이 렌더물에도 반영 (게이트 경유 역증명)
# ════════════════════════════════════════════════════════════════════
def test_disclaimer_injected_into_rendered_pptx(tmp_path):
    spy = DocSpy()
    # 면책 없는 깨끗한 데이터 → 게이트가 자동주입 후 렌더
    res = BUILD_PPTX.run(_clean_data(), spy, level="L1", out_dir=str(tmp_path))
    assert res.delivered is True
    # 렌더된 pptx 텍스트에 면책 포함 확인 (markitdown 없이 pptx 직접 파싱)
    from pptx import Presentation

    prs = Presentation(res.render.path)
    all_text = "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert REQUIRED_DISCLAIMER in all_text


# ════════════════════════════════════════════════════════════════════
# 4) as_document 전달 (게이트 통과 후, 원본 경로 그대로)
# ════════════════════════════════════════════════════════════════════
def test_as_document_delivers_original_path(tmp_path):
    spy = DocSpy()
    res = BUILD_PPTX.run(_clean_data(), spy, level="L1", out_dir=str(tmp_path))
    assert spy.called is True
    # 전달된 경로가 렌더 산출물 경로와 동일 (재압축/변경 없음)
    assert spy.deliveries[0]["path"] == res.render.path
    assert spy.deliveries[0]["fmt"] == "pptx"
    assert spy.deliveries[0]["level"] == "L1"


# ════════════════════════════════════════════════════════════════════
# 5) 비전 디자인 검수 (점수 산출)
# ════════════════════════════════════════════════════════════════════
def test_design_review_scores(tmp_path):
    spy = DocSpy()
    res = BUILD_PPTX.run(_clean_data(), spy, level="L1", out_dir=str(tmp_path))
    review = review_design(res.render, text=res.produced.text)
    assert set(review.breakdown) == set(DESIGN_RUBRIC)
    assert 0 <= review.total <= 100
    assert isinstance(review.passed, bool)
    # 렌더러가 브랜딩/정렬 보장 → 해당 항목 만점
    assert review.breakdown["branding"] == DESIGN_RUBRIC["branding"]


def test_design_review_vision_fn_injection(tmp_path):
    """vision_fn 주입(B-4 연결) 시 점수 보정 + vision_used=True."""
    spy = DocSpy()
    res = BUILD_PDF.run(_clean_data(), spy, level="L1", out_dir=str(tmp_path))

    def fake_vision(render):
        return {"breakdown": {"readability": 25}, "notes": ["vision: 가독성 우수"]}

    review = review_design(res.render, text=res.produced.text, vision_fn=fake_vision)
    assert review.vision_used is True
    assert review.breakdown["readability"] == 25
    assert any("vision" in n for n in review.notes)


def test_design_review_vision_fn_failure_graceful(tmp_path):
    spy = DocSpy()
    res = BUILD_PDF.run(_clean_data(), spy, level="L1", out_dir=str(tmp_path))

    def broken_vision(render):
        raise RuntimeError("vision API down")

    review = review_design(res.render, text=res.produced.text, vision_fn=broken_vision)
    # 비전 실패해도 휴리스틱으로 점수 산출 (우아 강등)
    assert review.total > 0
    assert any("heuristic only" in n for n in review.notes)


# ════════════════════════════════════════════════════════════════════
# 6) split_sections 단위 (렌더 입력 분해)
# ════════════════════════════════════════════════════════════════════
def test_split_sections_basic():
    # 새 파서: '# 제목'은 표지 타이틀, '##'는 섹션 heading.
    secs = dr.split_sections("# 표지\n## 소제목1\n본문1\n## 소제목2\n본문2")
    heads = [h for h, _ in secs]
    assert "소제목1" in heads and "소제목2" in heads
    bodies = {h: b for h, b in secs}
    assert "본문1" in bodies["소제목1"] and "본문2" in bodies["소제목2"]
